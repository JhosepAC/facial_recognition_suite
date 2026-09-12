"""
Tests del módulo de Video (Fase 2).

No dependen de insightface/PySide6: RecognitionService.recognize_frame se
sustituye por un doble de prueba (monkeypatch) para validar la orquestación
de VideoService sin requerir los modelos de IA reales.
"""
from pathlib import Path

import cv2
import numpy as np
import pytest

from app.database.base import Base
from app.database.models import Person
from app.database.repositories.person_repository import PersonRepository
from app.recognition.recognition_service import RecognitionService, RecognizedFace
from app.services.video_service import VideoService
from app.vision.video_processor import (
    VideoReader, resize_for_detection, scale_bbox, crop_with_padding,
)


@pytest.fixture()
def sample_video(tmp_path) -> str:
    """Genera un video sintético corto (60 frames, 30fps) para pruebas."""
    path = str(tmp_path / "sample.mp4")
    w, h, fps, n_frames = 640, 480, 30, 60
    writer = cv2.VideoWriter(path, cv2.VideoWriter_fourcc(*"mp4v"), fps, (w, h))
    for i in range(n_frames):
        frame = np.full((h, w, 3), 30, dtype=np.uint8)
        cv2.rectangle(frame, (50 + i, 100), (150 + i, 250), (180, 160, 140), -1)
        writer.write(frame)
    writer.release()
    return path


@pytest.fixture()
def session(tmp_path):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker

    engine = create_engine(f"sqlite:///{tmp_path / 'test.db'}")
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine)
    s = SessionLocal()
    yield s
    s.close()


# ------------------------------------------------------------------ #
# video_processor.py
# ------------------------------------------------------------------ #
def test_video_reader_metadata(sample_video):
    with VideoReader(sample_video) as reader:
        meta = reader.open()
        assert meta.total_frames == 60
        assert meta.fps == pytest.approx(30.0, abs=0.5)
        assert meta.width == 640 and meta.height == 480


def test_video_reader_sampling(sample_video):
    with VideoReader(sample_video) as reader:
        reader.open()
        samples = list(reader.iter_sampled_frames(interval_frames=10))
        assert [s.frame_number for s in samples] == [0, 10, 20, 30, 40, 50]


def test_video_reader_random_access(sample_video):
    with VideoReader(sample_video) as reader:
        reader.open()
        frame = reader.read_frame_at(30)
        assert frame is not None
        assert frame.shape == (480, 640, 3)


def test_resize_and_scale_bbox_roundtrip():
    frame = np.zeros((720, 1280, 3), dtype=np.uint8)
    resized, scale = resize_for_detection(frame, max_width=640)
    assert resized.shape[1] == 640
    assert scale == pytest.approx(0.5)

    bbox_resized = (100.0, 100.0, 200.0, 200.0)
    bbox_original = scale_bbox(bbox_resized, scale)
    assert bbox_original == (200.0, 200.0, 400.0, 400.0)


def test_crop_with_padding_stays_within_bounds():
    frame = np.zeros((480, 640, 3), dtype=np.uint8)
    crop = crop_with_padding(frame, (0.0, 0.0, 50.0, 50.0), padding_ratio=1.0)
    assert crop.shape[0] <= 480 and crop.shape[1] <= 640
    assert crop.size > 0


# ------------------------------------------------------------------ #
# VideoService (orquestación completa)
# ------------------------------------------------------------------ #
def test_analyze_video_end_to_end(session, sample_video, monkeypatch, tmp_path):
    person = PersonRepository(session).add(
        Person(nombre="Test", apellidos="User", empresa="QA")
    )
    session.commit()

    call_count = {"n": 0}

    def fake_recognize_frame(self, image_bgr, person_lookup=None, log_event=False,
                              origen="video", usuario=None, log_only_matches=False):
        call_count["n"] += 1
        h, w = image_bgr.shape[:2]
        bbox = (w * 0.2, h * 0.2, w * 0.6, h * 0.6)
        if call_count["n"] % 2 == 0:
            return [RecognizedFace(bbox=bbox, person_uuid=person.uuid,
                                    person_nombre=person.nombre_completo,
                                    confidence_pct=88.0, distance=0.2)]
        return [RecognizedFace(bbox=bbox, person_uuid=None, person_nombre=None,
                                confidence_pct=0.0, distance=1.0)]

    monkeypatch.setattr(RecognitionService, "recognize_frame", fake_recognize_frame)

    from app.core.config import settings
    monkeypatch.setattr(settings.video, "evidence_dir", str(tmp_path / "evidence"))
    monkeypatch.setattr(settings.video, "sample_interval_frames", 10)

    service = VideoService(session)
    progress_calls = []
    job = service.analyze_video(
        sample_video,
        progress_callback=lambda cur, total: progress_calls.append((cur, total)),
        usuario="tester",
    )

    assert job.estado == "completado"
    assert job.total_frames == 60
    assert job.fps == pytest.approx(30.0, abs=0.5)
    assert len(progress_calls) > 0
    assert progress_calls[-1][0] == job.total_frames  # progreso final = 100%

    detections = service.list_detections(job.id)
    assert len(detections) == 6  # 60 frames / intervalo 10

    matched = [d for d in detections if d.person_uuid == person.uuid]
    unmatched = [d for d in detections if d.person_uuid is None]
    assert len(matched) == 3
    assert len(unmatched) == 3

    for d in detections:
        assert d.evidencia_path is not None
        assert Path(d.evidencia_path).exists()


def test_analyze_video_invalid_extension_raises(session, tmp_path):
    bad_file = tmp_path / "not_a_video.txt"
    bad_file.write_text("hola")

    service = VideoService(session)
    with pytest.raises(Exception):
        service.validate_video_file(str(bad_file))


def test_delete_job_removes_evidence_files(session, sample_video, monkeypatch, tmp_path):
    def fake_recognize_frame(self, image_bgr, person_lookup=None, log_event=False,
                              origen="video", usuario=None, log_only_matches=False):
        h, w = image_bgr.shape[:2]
        return [RecognizedFace(bbox=(10, 10, 50, 50), person_uuid=None,
                                person_nombre=None, confidence_pct=0.0, distance=1.0)]

    monkeypatch.setattr(RecognitionService, "recognize_frame", fake_recognize_frame)
    from app.core.config import settings
    monkeypatch.setattr(settings.video, "evidence_dir", str(tmp_path / "evidence2"))
    monkeypatch.setattr(settings.video, "sample_interval_frames", 20)

    service = VideoService(session)
    job = service.analyze_video(sample_video)
    detections = service.list_detections(job.id)
    paths = [Path(d.evidencia_path) for d in detections]
    assert all(p.exists() for p in paths)

    service.delete_job(job.id)
    session.commit()  # en producción esto lo hace get_session() al salir del "with"
    assert all(not p.exists() for p in paths)
    assert service.get_job(job.id) is None


def test_analyze_video_caps_evidence_files(session, sample_video, monkeypatch, tmp_path):
    """A4: videos largos no deben agotar el disco con evidencias ilimitadas."""
    def fake_recognize_frame(self, image_bgr, person_lookup=None, log_event=False,
                              origen="video", usuario=None, log_only_matches=False):
        h, w = image_bgr.shape[:2]
        return [RecognizedFace(bbox=(10, 10, 50, 50), person_uuid=None,
                                person_nombre=None, confidence_pct=0.0, distance=1.0)]

    monkeypatch.setattr(RecognitionService, "recognize_frame", fake_recognize_frame)
    from app.core.config import settings
    monkeypatch.setattr(settings.video, "evidence_dir", str(tmp_path / "evidence_cap"))
    monkeypatch.setattr(settings.video, "sample_interval_frames", 10)
    monkeypatch.setattr(settings.video, "max_evidence_per_job", 2)

    service = VideoService(session)
    job = service.analyze_video(sample_video)
    detections = service.list_detections(job.id)
    assert len(detections) == 6
    # Las detecciones siguen guardándose aunque se limite la evidencia en disco.
    with_evidence = [d for d in detections if d.evidencia_path is not None]
    assert len(with_evidence) == 2
    assert sum(1 for d in detections if d.evidencia_path is None) == 4


# ------------------------------------------------------------------ #
# Exportación de detecciones (CSV / Excel)
# ------------------------------------------------------------------ #
@pytest.fixture()
def completed_job(session, sample_video, monkeypatch, tmp_path):
    person = PersonRepository(session).add(
        Person(nombre="Export", apellidos="Tester", empresa="QA")
    )
    session.commit()

    def fake_recognize_frame(self, image_bgr, person_lookup=None, log_event=False,
                              origen="video", usuario=None, log_only_matches=False):
        h, w = image_bgr.shape[:2]
        bbox = (w * 0.2, h * 0.2, w * 0.6, h * 0.6)
        return [RecognizedFace(bbox=bbox, person_uuid=person.uuid,
                                person_nombre=person.nombre_completo,
                                confidence_pct=92.0, distance=0.15)]

    monkeypatch.setattr(RecognitionService, "recognize_frame", fake_recognize_frame)
    from app.core.config import settings
    monkeypatch.setattr(settings.video, "evidence_dir", str(tmp_path / "evidence3"))
    monkeypatch.setattr(settings.video, "sample_interval_frames", 10)

    service = VideoService(session)
    job = service.analyze_video(sample_video)
    assert job.estado == "completado"
    return service, job


def test_export_detections_csv(completed_job, tmp_path):
    service, job = completed_job
    dest = tmp_path / "detections.csv"
    count = service.export_detections(job.id, str(dest), fmt="csv", usuario="tester")
    assert count == 6
    assert dest.exists()

    import pandas as pd
    df = pd.read_csv(dest)
    assert list(df.columns) == ["frame", "tiempo", "persona", "confianza_pct",
                                "distancia_coseno", "conf_deteccion", "estado",
                                "evidencia", "gafas", "mascarilla", "barba",
                                "bigote", "sonrisa", "ojos_abiertos"]
    assert df["estado"].eq("Reconocido").all()


def test_export_detections_excel_styled(completed_job, tmp_path):
    service, job = completed_job
    dest = tmp_path / "detections.xlsx"
    count = service.export_detections(job.id, str(dest), fmt="excel", usuario="tester")
    assert count == 6
    assert dest.exists()

    from openpyxl import load_workbook
    wb = load_workbook(dest)
    assert wb.sheetnames == ["Resumen", "Detecciones"]

    ws_d = wb["Detecciones"]
    attr_headers = ["Gafas", "Mascarilla", "Barba", "Bigote", "Sonrisa", "Ojos abiertos"]
    headers = [ws_d.cell(row=1, column=c).value for c in range(1, 16)]
    assert headers == ["#", "Frame", "Tiempo (hh:mm:ss)", "Persona", "Confianza (%)",
                       "Conf. detección", "Distancia coseno", "Estado", "Evidencia",
                       *attr_headers]
    assert ws_d.freeze_panes == "A2"
    assert ws_d.auto_filter.ref is not None
    assert ws_d.cell(row=2, column=9).value  # evidencia presente
    assert ws_d.cell(row=2, column=10).value == "—"  # persona sin análisis facial

    ws_sum = wb["Resumen"]
    assert "Análisis de video" in ws_sum["A1"].value
    assert ws_sum["A1"].font.bold
