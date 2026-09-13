"""Video service layer. Orchestrates:
  VideoReader (frame extraction) + RecognitionService (detection/match)
  + VideoJobRepository (persistence) + evidence storage on disk.

The GUI (or any caller) should only use VideoService; it must not access
VideoReader or RecognitionService directly to keep the strict UI/logic
separation.
"""
from __future__ import annotations

import json
import uuid
from datetime import datetime
from pathlib import Path
from typing import Callable

import cv2
import pandas as pd

from app.core.config import settings
from app.core.exceptions import BioVisionError
from app.core.logger import audit_logger, logger
from app.database.models import Person, VideoDetection, VideoJob
from app.database.repositories.video_repository import VideoJobRepository
from app.recognition.recognition_service import RecognitionService
from app.services.export_service import slugify
from app.services.statistics_service import primary_embedding_attrs
from app.utils.spreadsheet import sanitize_formula
from app.vision.face_attributes import ATTR_FIELDS
from app.vision.video_processor import (
    VideoReader, crop_with_padding, resize_for_detection, scale_bbox,
)

_ATTR_HEADER_LABELS = {
    "gafas": "Gafas",
    "mascarilla": "Mascarilla",
    "barba": "Barba",
    "bigote": "Bigote",
    "sonrisa": "Sonrisa",
    "ojos_abiertos": "Ojos abiertos",
}

ProgressCallback = Callable[[int, int], None]  # (frames_procesados, total_frames) -> None


class VideoService:
    def __init__(self, session):
        self.session = session
        self.repo = VideoJobRepository(session)
        self.recognition = RecognitionService(session)
        self.evidence_dir = settings.resolve_path(settings.video.evidence_dir)
        self.evidence_dir.mkdir(parents=True, exist_ok=True)

    # ------------------------------------------------------------------ #
    def validate_video_file(self, file_path: str) -> None:
        path = Path(file_path)
        if not path.exists():
            raise BioVisionError(f"File does not exist: {file_path}")
        if path.suffix.lower() not in settings.video.allowed_extensions:
            allowed = ", ".join(settings.video.allowed_extensions)
            raise BioVisionError(
                f"Formato no soportado ({path.suffix}). Formatos permitidos: {allowed}"
            )

    # ------------------------------------------------------------------ #
    def analyze_video(
        self,
        file_path: str,
        progress_callback: ProgressCallback | None = None,
        usuario: str | None = None,
        sample_interval: int | None = None,
    ) -> VideoJob:
        """
        Analiza un video completo: extrae frames muestreados, detecta y reconoce
        rostros, guarda evidencia y persiste cada detección. Devuelve el VideoJob
        final con su estado ("completado" o "error").

        `sample_interval`: número de frames entre muestras (sobrescribe el valor
        configurado en `settings.video.sample_interval_frames`).
        """
        self.validate_video_file(file_path)
        path = Path(file_path)

        job = VideoJob(
            file_path=str(path),
            nombre_archivo=path.name,
            estado="procesando",
            usuario=usuario,
        )
        self.repo.add_job(job)
        self.session.commit()  # visible de inmediato para la GUI (lista de trabajos)

        person_lookup = {p.uuid: p.nombre_completo for p in self.session.query(Person).all()}

        reader = VideoReader(str(path))
        try:
            meta = reader.open()
            job.fps = meta.fps
            job.total_frames = meta.total_frames
            job.duracion_seg = meta.duration_seg
            self.session.commit()

            interval = sample_interval or settings.video.sample_interval_frames
            max_width = settings.video.max_processing_width
            padding = settings.video.evidence_crop_padding
            save_evidence = settings.video.save_evidence_thumbnails

            processed_samples = 0
            evidence_saved = 0
            evidence_cap = settings.video.max_evidence_per_job
            for sample in reader.iter_sampled_frames(interval):
                detection_frame, scale = resize_for_detection(sample.frame_bgr, max_width)

                try:
                    recognized_faces = self.recognition.recognize_frame(
                        detection_frame,
                        person_lookup=person_lookup,
                        log_event=True,
                        origen="video",
                        usuario=usuario,
                        log_only_matches=True,
                    )
                except Exception as exc:  # noqa: BLE001
                    logger.warning("Frame {} omitido por error de reconocimiento: {}",
                                   sample.frame_number, exc)
                    recognized_faces = []

                for face in recognized_faces:
                    bbox_original = scale_bbox(face.bbox, scale)

                    evidence_path = None
                    if save_evidence and evidence_saved < evidence_cap:
                        crop = crop_with_padding(sample.frame_bgr, bbox_original, padding)
                        if crop.size > 0:
                            video_stem = (Path(job.file_path).stem
                                          if job.file_path else f"job{job.id}")
                            persona_slug = slugify(face.person_nombre or "desconocido",
                                                   "desconocido")
                            fecha = datetime.now().strftime("%Y%m%d_%H%M%S")
                            fname = (f"evidencia_{uuid.uuid4().hex[:4]}_{video_stem}_"
                                     f"{persona_slug}_f{sample.frame_number}_{fecha}.jpg")
                            dest = self.evidence_dir / fname
                            cv2.imwrite(str(dest), crop)
                            evidence_path = str(dest)
                            evidence_saved += 1

                    detection = VideoDetection(
                        video_job_id=job.id,
                        frame_number=sample.frame_number,
                        timestamp_seg=sample.timestamp_seg,
                        person_uuid=face.person_uuid,
                        confianza=face.confidence_pct,
                        distancia=face.distance,
                        det_confidence=face.det_score,
                        facial_attributes=json.dumps(face.attributes.to_dict(), ensure_ascii=False)
                        if getattr(face, "attributes", None) else None,
                        bbox_x1=bbox_original[0], bbox_y1=bbox_original[1],
                        bbox_x2=bbox_original[2], bbox_y2=bbox_original[3],
                        evidencia_path=evidence_path,
                    )
                    self.repo.add_detection(detection)

                processed_samples += 1
                job.frames_procesados = sample.frame_number + 1
                if processed_samples % 5 == 0:
                    self.session.commit()
                    if progress_callback:
                        progress_callback(job.frames_procesados, job.total_frames or 1)

            job.estado = "completado"
            job.fecha_finalizacion = datetime.utcnow()
            self.session.commit()

            if progress_callback:
                progress_callback(job.total_frames or 1, job.total_frames or 1)

            audit_logger.info(
                "Video analizado | job_id={} | archivo={} | detecciones={} | usuario={}",
                job.id, job.nombre_archivo, self.repo.count_detections(job.id), usuario or "sistema",
            )

        except Exception as exc:  # noqa: BLE001
            logger.exception("Error analizando video {}", file_path)
            job.estado = "error"
            job.mensaje_error = str(exc)
            job.fecha_finalizacion = datetime.utcnow()
            self.session.commit()
        finally:
            reader.close()

        return job

    # ------------------------------------------------------------------ #
    def list_jobs(self, limit: int = 100) -> list[VideoJob]:
        return list(self.repo.list_jobs(limit=limit))

    def get_job(self, job_id: int) -> VideoJob | None:
        return self.repo.get_job(job_id)

    def list_detections(self, job_id: int) -> list[VideoDetection]:
        return list(self.repo.list_detections(job_id))

    def job_summary(self, job_id: int) -> dict:
        """Resumen del análisis: totales y personas detectadas con su miniatura."""
        detections = self.list_detections(job_id)
        total = len(detections)
        recognized = sum(1 for d in detections if d.person_uuid is not None)

        persons: dict[str, dict] = {}
        for d in detections:
            if d.person_uuid is None:
                continue
            person = d.person
            entry = persons.get(d.person_uuid)
            if entry is None:
                thumb = None
                if person and person.photos:
                    primary = next((p for p in person.photos if p.es_principal), person.photos[0])
                    thumb = primary.thumbnail_path
                entry = {
                    "uuid": d.person_uuid,
                    "nombre": person.nombre_completo if person else "Desconocido",
                    "thumbnail": thumb,
                    "count": 0,
                    "confianzas": [],
                }
                persons[d.person_uuid] = entry
            entry["count"] += 1
            if d.confianza is not None:
                entry["confianzas"].append(d.confianza)

        for entry in persons.values():
            conf = entry.pop("confianzas", [])
            entry["confianza_promedio"] = round(sum(conf) / len(conf), 1) if conf else None

        return {
            "total": total,
            "reconocidas": recognized,
            "no_reconocidas": total - recognized,
            "tasa_reconocimiento_pct": round((recognized / total) * 100, 1) if total else 0.0,
            "personas": sorted(persons.values(), key=lambda e: e["count"], reverse=True),
        }

    def delete_job(self, job_id: int, usuario: str | None = None) -> bool:
        job = self.repo.get_job(job_id)
        if job is None:
            return False
        # 1) DB is deleted first and committed: if it fails, evidence files
        #    remain on disk and are not irreversibly lost.
        evidence_paths = [d.evidencia_path for d in job.detections if d.evidencia_path]
        deleted = self.repo.delete_job(job_id)
        self.session.commit()
        # 2) Only then are evidence files removed (file deletion failures
        #    must not abort the DB operation).
        for path in evidence_paths:
            try:
                Path(path).unlink(missing_ok=True)
            except OSError as exc:
                logger.warning("No se pudo eliminar evidencia '{}': {}", path, exc)
        audit_logger.info("Trabajo de video eliminado | job_id={} | usuario={}",
                           job_id, usuario or "sistema")
        return deleted

    # ------------------------------------------------------------------ #
    # Detection export (CSV / Excel with styled output).
    # ------------------------------------------------------------------ #
    def export_detections(
        self,
        job_id: int,
        path: str,
        fmt: str = "csv",
        usuario: str | None = None,
    ) -> int:
        """Export job detections to CSV or Excel.

        Args:
            job_id: Identifier of the video job.
            path: Destination file path.
            fmt: Export format (``csv`` or ``excel``).
            usuario: Acting username for audit logging.

        Returns:
            Number of exported records.
        """
        job = self.get_job(job_id)
        if job is None:
            raise BioVisionError("Video analysis does not exist.")
        detections = self.list_detections(job_id)
        fmt = (fmt or "csv").lower()

        if fmt == "csv":
            self._export_detections_csv(job, detections, path)
        elif fmt in ("excel", "xlsx"):
            self._export_detections_excel(job, detections, path)
        else:
            raise BioVisionError("Formato no soportado. Usa 'csv' o 'excel'.")

        audit_logger.info(
            "Detecciones de video exportadas | job_id={} | formato={} | detecciones={} | destino={} | usuario={}",
            job_id, fmt, len(detections), Path(path).name, usuario or "sistema",
        )
        return len(detections)

    @staticmethod
    def _detections_rows(job: VideoJob, detections) -> list[dict]:
        rows = []
        for d in detections:
            attrs = primary_embedding_attrs(d.person) if d.person is not None else None
            row = {
                "frame": d.frame_number,
                "tiempo": d.timestamp_fmt,
                "persona": sanitize_formula(d.person.nombre_completo) if d.person else "Desconocido",
                "confianza_pct": d.confianza,
                "distancia_coseno": d.distancia,
                "conf_deteccion": d.det_confidence,
                "estado": "Reconocido" if d.person_uuid is not None else "Desconocido",
                "evidencia": sanitize_formula(d.evidencia_path or ""),
            }
            for field in ATTR_FIELDS:
                row[field] = bool(getattr(attrs, field, False)) if attrs else None
            rows.append(row)
        return rows

    def _export_detections_csv(self, job: VideoJob, detections, path: str) -> None:
        df = pd.DataFrame(self._detections_rows(job, detections))
        df.to_csv(path, index=False, encoding="utf-8-sig")

    def _export_detections_excel(self, job: VideoJob, detections, path: str) -> None:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        BLUE = "2F6FED"
        BLUE_DARK = "17408F"
        GREEN = "2E7D32"
        GREEN_SOFT = "E8F5E9"
        RED = "C0392B"
        RED_SOFT = "FDECEA"
        GREY_SOFT = "F5F6FA"
        TEXT_MUTED = "6B6E7D"

        wb = Workbook()

        # ---- Sheet 1: Summary ----
        ws = wb.active
        ws.title = "Resumen"
        ws.sheet_view.showGridLines = False

        summary = self.job_summary(job.id)

        ws.merge_cells("A1:F1")
        cell = ws["A1"]
        cell.value = sanitize_formula(f"Análisis de video — {job.nombre_archivo}")
        cell.font = Font(size=16, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 30

        ws.merge_cells("A2:F2")
        ws["A2"] = f"Generado el {datetime.now():%Y-%m-%d %H:%M:%S}  ·  Tiempo de navegación en frames muestreados"
        ws["A2"].font = Font(size=9, color=TEXT_MUTED)

        kpi_rows = [
            ("Archivo", sanitize_formula(job.nombre_archivo)),
            ("Detecciones totales", summary["total"]),
            ("Rostros reconocidos", summary["reconocidas"]),
            ("Rostros no reconocidos", summary["no_reconocidas"]),
            ("Tasa de reconocimiento", f"{summary['tasa_reconocimiento_pct']:.1f}%"),
            ("Personas únicas", len(summary["personas"])),
            ("Duración (seg)", f"{job.duracion_seg:.1f}" if job.duracion_seg else "—"),
            ("Frames procesados", f"{job.frames_procesados or 0} / {job.total_frames or 0}"),
            ("FPS", f"{job.fps:.1f}" if job.fps else "—"),
        ]

        start = 4
        for i, (label, value) in enumerate(kpi_rows):
            r = start + i
            ws.cell(row=r, column=1, value=label).font = Font(bold=True, color=TEXT_MUTED)
            ws.cell(row=r, column=1).alignment = Alignment(horizontal="right")
            ws.cell(row=r, column=2, value=value).font = Font(bold=True, color="1D1E24")
            if i % 2 == 1:
                ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=GREY_SOFT)
                ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=GREY_SOFT)

        # Persons detected in the summary
        pr = start + len(kpi_rows) + 2
        ws.cell(row=pr, column=1, value="Personas detectadas").font = Font(size=12, bold=True, color=BLUE_DARK)
        pr += 1
        headers_pers = ["Persona", "Apariciones", "Confianza promedio"]
        for c, h in enumerate(headers_pers, start=1):
            ccell = ws.cell(row=pr, column=c, value=h)
            ccell.font = Font(bold=True, color="FFFFFF")
            ccell.fill = PatternFill("solid", fgColor=BLUE_DARK)
        pr += 1
        for entry in summary["personas"][:20]:
            ws.cell(row=pr, column=1, value=sanitize_formula(entry["nombre"]))
            ws.cell(row=pr, column=2, value=entry["count"])
            ws.cell(row=pr, column=3, value=(f"{entry['confianza_promedio']:.1f}" if entry.get("confianza_promedio") is not None else "—"))
            pr += 1

        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 22
        ws.column_dimensions["C"].width = 22

        # ---- Sheet "Detections" ----
        ws_d = wb.create_sheet("Detecciones")
        attr_headers = [_ATTR_HEADER_LABELS.get(f, f) for f in ATTR_FIELDS]
        headers = ["#", "Frame", "Tiempo (hh:mm:ss)", "Persona", "Confianza (%)",
                   "Conf. detección", "Distancia coseno", "Estado", "Evidencia",
                   *attr_headers]
        header_fill = PatternFill("solid", fgColor=BLUE)
        header_font = Font(bold=True, color="FFFFFF", size=11)
        thin = Side(style="thin", color="D9DCE4")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for c, h in enumerate(headers, start=1):
            cell = ws_d.cell(row=1, column=c, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        ws_d.row_dimensions[1].height = 22

        for i, det in enumerate(detections, start=1):
            r = i + 1
            recognized = det.person_uuid is not None
            attrs = primary_embedding_attrs(det.person) if det.person is not None else None
            values = [
                i,
                det.frame_number,
                det.timestamp_fmt,
                sanitize_formula(det.person.nombre_completo) if det.person else "Desconocido",
                round(det.confianza, 1) if det.confianza is not None else "—",
                round(det.det_confidence, 2) if det.det_confidence is not None else "—",
                round(det.distancia, 4) if det.distancia is not None else "—",
                "Reconocido" if recognized else "Desconocido",
                sanitize_formula(det.evidencia_path or ""),
                *(("Sí" if getattr(attrs, f, False) else "No") if attrs else "—"
                  for f in ATTR_FIELDS),
            ]
            for c, v in enumerate(values, start=1):
                cell = ws_d.cell(row=r, column=c, value=v)
                cell.border = border
                if c in (1, 3, 5, 6, 7, 8):
                    cell.alignment = Alignment(horizontal="center")
                if c == 8:
                    cell.font = Font(bold=True, color=(GREEN if recognized else RED))
                    cell.fill = PatternFill("solid", fgColor=(GREEN_SOFT if recognized else RED_SOFT))
                elif i % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor=GREY_SOFT)

        width_map = [6, 8, 18, 30, 15, 14, 16, 13, 46] + [12] * len(ATTR_FIELDS)
        for idx, w in enumerate(width_map, start=1):
            ws_d.column_dimensions[get_column_letter(idx)].width = w

        ws_d.freeze_panes = "A2"
        ws_d.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws_d.max_row}"
        ws_d.sheet_view.showGridLines = False

        wb.save(path)
